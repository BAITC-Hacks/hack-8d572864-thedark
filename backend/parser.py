"""Extract original text and stable source locations without executing document content."""
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 200_000
MAX_BLOCKS = 2000


def extract(filename: str, content: bytes, side: str) -> list[dict]:
    filename = filename.replace('\\', '/').rsplit('/', 1)[-1]
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.docx', '.pdf', '.xlsx'}:
        raise ValueError('PDF, DOCX немесе XLSX файлы қажет.')
    if not content or len(content) > MAX_BYTES:
        raise ValueError('Бос файл немесе файл көлемі 20 MB-тан артық.')
    if suffix in {'.docx', '.xlsx'}:
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 80 * 1024 * 1024:
                    raise ValueError('Құжаттың ашылған көлемі тым үлкен.')
        except BadZipFile as exc:
            raise ValueError('DOCX/XLSX файлының ішкі құрылымы бұзылған.') from exc
    blocks, current_point, characters = [], '', 0

    def add(text: str, location: str):
        nonlocal current_point, characters
        text = text.strip()
        if not text:
            return
        characters += len(text)
        if characters > MAX_CHARS or len(blocks) >= MAX_BLOCKS:
            raise ValueError('Құжат тым ұзын: MVP шегі 200 000 таңба / 2 000 үзінді. Құжатты бөліңіз.')
        match = re.match(r'^\s*(\d+(?:\.\d+)*)(?:[.)]|\s)', text)
        if match:
            current_point = match.group(1)
        prefix = 'A' if side == 'before' else 'B'
        blocks.append({'id': f'{prefix}{len(blocks)+1:04}', 'document': filename,
                       'filename': filename, 'side': side, 'point': current_point or location,
                       'location': location, 'text': text})

    if suffix == '.docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document = Document(BytesIO(content))
        paragraph_number, table_number = 0, 0
        for child in document.element.body.iterchildren():
            if child.tag == qn('w:p'):
                paragraph_number += 1
                add(Paragraph(child, document).text, f'Абзац {paragraph_number}')
            elif child.tag == qn('w:tbl'):
                table_number += 1
                for row_number, row in enumerate(Table(child, document).rows, 1):
                    add(' | '.join(cell.text for cell in row.cells), f'Кесте {table_number}, жол {row_number}')
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('Құпиясөзбен қорғалған PDF қолдау таппайды.')
        if len(reader.pages) > 300:
            raise ValueError('MVP шегі: PDF 300 беттен аспауы керек.')
        for page_number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip() and list(page.images):
                raise ValueError(f'PDF {page_number}-беті скан болуы мүмкін. Алдымен OCR қажет.')
            for line_number, line in enumerate(text.splitlines(), 1):
                add(line, f'Бет {page_number}, жол {line_number}')
    else:
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            visited = 0
            for sheet in book:
                current_point = ''
                if sheet.max_column and sheet.max_column > 200:
                    raise ValueError('XLSX парағында 200-ден артық баған бар; керекті бағандарды бөліңіз.')
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    visited += 1
                    if visited > 10000:
                        raise ValueError('MVP шегі: XLSX 10 000 жолдан аспауы керек.')
                    if any(value is not None and str(value).strip() for value in row):
                        add(' | '.join('' if value is None else str(value) for value in row), f'{sheet.title}!{row_number}')
        finally:
            book.close()
    if not blocks:
        raise ValueError('Оқылатын мәтін табылмады. Скан құжат үшін OCR қажет.')
    return blocks
